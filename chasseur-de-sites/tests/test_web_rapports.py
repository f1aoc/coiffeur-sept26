"""Routes du Lot 4 : rapport PDF, ZIP groupé, exports, agence, modèles de messages."""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import load_workbook
from PIL import Image

from conftest import moteurs_pdf_disponibles, texte_pdf


@pytest.fixture
def web_rapports(web, base_rapports, monkeypatch):
    moteurs = moteurs_pdf_disponibles()
    if moteurs:
        monkeypatch.setenv("CHASSEUR_MOTEUR_PDF", moteurs[0])
    return web, base_rapports, bool(moteurs)


def test_rapport_pdf_depuis_la_fiche(web_rapports):
    web, ids, moteur = web_rapports
    fiche = web.get(f"/prospects/{ids['casse']}").text
    assert f'href="/prospects/{ids["casse"]}/rapport.pdf"' in fiche and "Générer le rapport PDF" in fiche
    assert "Votre site présente 6 problèmes qui font fuir vos visiteurs." in fiche
    if not moteur:
        pytest.skip("aucun moteur PDF")
    r = web.get(f"/prospects/{ids['casse']}/rapport.pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert "diagnostic-salon-casse.pdf" in r.headers["content-disposition"]
    pages, texte = texte_pdf(r.content)
    assert pages == 2 and "Atelier Web Provence" in texte
    assert web.get("/prospects/999/rapport.pdf").status_code == 404


def test_rapports_groupes_en_zip(web_rapports):
    web, ids, moteur = web_rapports
    page = web.get("/resultats").text
    assert page.count('name="ids"') == 3 and 'form="selection"' in page and 'id="tout-cocher"' in page
    assert web.post("/rapports.zip", data={}).status_code == 400
    if not moteur:
        pytest.skip("aucun moteur PDF")
    r = web.post("/rapports.zip", data={"ids": [ids["casse"], ids["correct"], ids["casse"]]})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        assert sorted(z.namelist()) == ["diagnostic-coiffure-impeccable.pdf", "diagnostic-salon-casse.pdf"]


def test_exports_de_la_liste_filtree(web_rapports):
    web, ids, _ = web_rapports
    page = web.get("/resultats?etat=Obsolète").text
    import re

    lien = re.search(r'href="(/export\.xlsx\?[^"]*)"', page)[1]
    assert "etat=Obsol%C3%A8te" in lien and "page=" not in lien
    r = web.get("/export.xlsx?etat=Obsolète")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    feuille = load_workbook(io.BytesIO(r.content))["Prospects"]
    assert [row[0] for row in feuille.iter_rows(min_row=2, values_only=True)] == ["Institut Vieillot"]
    r = web.get("/export.csv?q=salon")
    assert r.status_code == 200 and "Salon Cassé" in r.content.decode("utf-8-sig")
    assert "Institut Vieillot" not in r.content.decode("utf-8-sig")


def test_messages_pre_remplis_sur_la_fiche(web_rapports):
    web, ids, _ = web_rapports
    fiche = web.get(f"/prospects/{ids['casse']}").text
    assert "Premier contact" in fiche and fiche.count('data-copier="message-') == 3
    assert "En consultant le site de Salon Cassé, j&#39;ai remarqué que votre adresse web affiche une page « domaine à vendre »" in fiche
    assert "Atelier Web Provence" in fiche and "{entreprise}" not in fiche
    assert "Répondez STOP" in fiche


def test_reglages_agence_logo_et_modeles(web_rapports):
    web, ids, _ = web_rapports
    page = web.get("/reglages").text
    assert 'value="Atelier Web Provence"' in page and 'value="#0a7f5a"' in page and "/agence/logo" in page
    assert web.get("/agence/logo").headers["content-security-policy"] == "script-src 'none'"

    png = io.BytesIO()
    Image.new("RGB", (120, 40), "navy").save(png, "PNG")
    r = web.post(
        "/reglages",
        data={"agence_nom": "Studio Pixel", "agence_couleur": "#ff6600", "agence_email": "contact@pixel.test",
              "agence_telephone": "", "agence_site": "", "agence_appel": "Parlons-en.",
              "modele_sms": "Bonjour {entreprise}, ici {agence}.", "parallelisme": "10"},
        files={"agence_logo": ("logo.png", png.getvalue(), "image/png")},
    )
    assert r.status_code == 303 and r.headers["location"] == "/reglages?ok=1"
    assert web.get("/agence/logo").headers["content-type"] == "image/png"
    fiche = web.get(f"/prospects/{ids['correct']}").text
    assert "Bonjour Coiffure Impeccable, ici Studio Pixel." in fiche
    html = web.get(f"/prospects/{ids['casse']}/rapport.html").text
    assert "Studio Pixel" in html and "#ff6600" in html and "Parlons-en." in html and "data:image/png;base64," in html


@pytest.mark.parametrize(
    "donnees, fichiers, message",
    [
        ({"agence_couleur": "rouge"}, None, "Couleur invalide"),
        ({"agence_email": "pas-un-email"}, None, "e-mail"),
        ({}, {"agence_logo": ("logo.svg", b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', "image/svg+xml")}, "refusé par sécurité"),
        ({}, {"agence_logo": ("logo.svg", b'<svg onload="alert(1)"/>', "image/svg+xml")}, "refusé par sécurité"),
        ({}, {"agence_logo": ("logo.gif", b"GIF89a", "image/gif")}, "PNG ou SVG"),
        ({}, {"agence_logo": ("logo.png", b"pas une image", "image/png")}, "PNG valide"),
    ],
)
def test_reglages_agence_refuses(web_rapports, donnees, fichiers, message):
    web, _, _ = web_rapports
    r = web.post("/reglages", data={"parallelisme": "10", **donnees}, files=fichiers)
    assert r.status_code == 303 and "erreur=" in r.headers["location"]
    page = web.get(r.headers["location"]).text
    assert message in page
    assert 'value="Atelier Web Provence"' in page  # rien n'a été modifié


def test_suppression_du_logo(web_rapports):
    web, _, _ = web_rapports
    web.post("/reglages", data={"supprimer_logo": "1", "parallelisme": "10"})
    assert web.get("/agence/logo").status_code == 404
