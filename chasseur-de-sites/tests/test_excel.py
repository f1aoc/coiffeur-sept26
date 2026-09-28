"""Export Excel et CSV de la liste filtrée."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from chasseur.controles import CONTROLES
from chasseur.db import depot
from chasseur.reports.excel import COLONNES, exporter_csv, exporter_excel


def classeur(stockage, **filtres):
    with stockage.session() as s:
        return load_workbook(io.BytesIO(exporter_excel(stockage, s, depot.Filtres(**filtres), "test")))


def test_colonnes_une_par_controle():
    assert COLONNES[:6] == ["Nom", "Ville", "Catégorie", "Score", "Statut", "Problèmes principaux"]
    for c in CONTROLES:
        if c.id != "url":
            assert c.libelle in COLONNES
    assert COLONNES[-2:] == ["Capture ordinateur", "Capture téléphone"]
    assert not any("clé" in c.lower() or "api" in c.lower() for c in COLONNES)


def test_feuille_prospects(stockage, base_rapports):
    wb = classeur(stockage)
    assert wb.sheetnames == ["Prospects", "Récapitulatif"]
    f = wb["Prospects"]
    assert [c.value for c in f[1]] == COLONNES
    assert f.freeze_panes == "B2"  # en-têtes figés
    assert f.auto_filter.ref == f"A1:{f.cell(1, len(COLONNES)).column_letter}4"
    lignes = {r[0]: dict(zip(COLONNES, r)) for r in f.iter_rows(min_row=2, values_only=True)}
    casse = lignes["Salon Cassé"]
    assert (casse["Catégorie"], casse["Score"], casse["Statut"], casse["Ville"]) == ("Cassé", 100, "À contacter", "Avignon")
    assert (casse["Note Google"], casse["Avis Google"]) == (4.6, 52)
    assert casse["Problèmes principaux"] == "Domaine parqué (à vendre), Page blanche, Pas de HTTPS"
    assert casse["Domaine expiré ou parqué"] == "Problème : Domaine parqué (à vendre)"
    assert casse["Domaine qui résout (DNS)"] == "OK"
    assert casse["Capture ordinateur"].endswith("0-bureau.webp")
    ancien = lignes["Institut Vieillot"]
    assert ancien["CMS ou technologies périmées"] == "Problème : WordPress périmé, jQuery périmé"
    assert lignes["Coiffure Impeccable"]["Catégorie"] == "Correct"
    # Ordre de l'écran Résultats (score décroissant)
    assert [r[0] for r in f.iter_rows(min_row=2, values_only=True)] == ["Salon Cassé", "Institut Vieillot", "Coiffure Impeccable"]


def test_couleurs_par_categorie(stockage, base_rapports):
    f = classeur(stockage)["Prospects"]
    col = COLONNES.index("Catégorie") + 1
    couleurs = {f.cell(r, 1).value: f.cell(r, col).fill.fgColor.rgb[-6:] for r in range(2, 5)}
    assert couleurs == {"Salon Cassé": "FDECEA", "Institut Vieillot": "FFF4E0", "Coiffure Impeccable": "E6F4EA"}
    assert f.cell(2, COLONNES.index("Domaine expiré ou parqué") + 1).font.color.rgb[-6:] == "B42318"


def test_recapitulatif(stockage, base_rapports):
    recap = classeur(stockage)["Récapitulatif"]
    valeurs = [tuple(r) for r in recap.iter_rows(values_only=True)]
    assert ("Par catégorie", "Prospects") in valeurs
    assert {("Cassé", 1), ("Obsolète", 1), ("Correct", 1)} <= set(valeurs)
    assert ("Site responsive", 2) in valeurs and ("HTTPS et redirection HTTP → HTTPS", 2) in valeurs
    assert ("Domaine expiré ou parqué", 1) in valeurs
    assert ("À contacter", 3) in valeurs
    assert "3 prospect(s)" in valeurs[1][0] and "test" in valeurs[1][0]


def test_export_respecte_les_filtres(stockage, base_rapports):
    f = classeur(stockage, etat="Obsolète")["Prospects"]
    assert [r[0] for r in f.iter_rows(min_row=2, values_only=True)] == ["Institut Vieillot"]
    f = classeur(stockage, probleme="responsive", tri="nom", ordre="asc")["Prospects"]
    assert [r[0] for r in f.iter_rows(min_row=2, values_only=True)] == ["Institut Vieillot", "Salon Cassé"]
    assert classeur(stockage, q="introuvable")["Prospects"].max_row == 1  # en-têtes seuls


def test_export_csv(stockage, base_rapports):
    with stockage.session() as s:
        brut = exporter_csv(stockage, s, depot.Filtres(statut="À contacter"))
    assert brut.startswith(b"\xef\xbb\xbf")
    lignes = list(csv.DictReader(brut.decode("utf-8-sig").splitlines(), delimiter=";"))
    assert list(lignes[0]) == COLONNES and len(lignes) == 3
    assert lignes[0]["Nom"] == "Salon Cassé" and lignes[2]["Score"] == "0"
