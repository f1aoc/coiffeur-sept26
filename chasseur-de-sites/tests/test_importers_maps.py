"""Import des fichiers de l'outil Google Maps existant (CSV et Excel), sans retouche."""

from __future__ import annotations

import pytest

from chasseur.importers import FORMAT_GENERIQUE, FORMAT_MAPS_CSV, FORMAT_MAPS_EXCEL, lire_prospects
from chasseur.importers.fichiers import detecter_format, ville_depuis_adresse

# Exactement ce qu'écrit find_no_website.py (DictWriter, virgules, UTF-8 avec BOM)
CSV_MAPS = (
    "﻿name,status,website,phone,address,category,rating,reviews,maps_url,place_id,query\r\n"
    "Salon Léa,none,,04 90 00 00 01,\"12 Rue de la République, 84800 L'Isle-sur-la-Sorgue, France\",Coiffeur,4.7,132,"
    "https://maps.google.com/?cid=1,ChIJ1,coiffeur L'Isle-sur-la-Sorgue\r\n"
    "Barber Max,social_or_directory:facebook.com,https://www.facebook.com/barbermax,04 90 00 00 02,"
    "\"3 Place Carnot, 84300 Cavaillon, France\",Barbier,4.2,18,https://maps.google.com/?cid=2,ChIJ2,coiffeur Cavaillon\r\n"
    "Coiffure Réelle,real,https://coiffure-reelle.fr/,04 90 00 00 03,\"1 Av. Foch, 84000 Avignon, France\",Coiffeur,,,"
    "https://maps.google.com/?cid=3,ChIJ3,coiffeur Avignon\r\n"
)


def ecrire_xlsx_maps(chemin):
    """Reproduit write_xlsx() de l'outil Google Maps (en-têtes français, lien sur « Ouvrir »)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    ws.append(["Nom", "Site web", "Page actuelle", "Téléphone", "Adresse", "Catégorie", "Note", "Avis", "Google Maps", "Recherche"])
    ws.append(["Salon Léa", "Aucun", "", "04 90 00 00 01", "12 Rue X, 84800 L'Isle-sur-la-Sorgue, France", "Coiffeur", 4.7, 132, "Ouvrir", "coiffeur"])
    ws.cell(row=2, column=9).hyperlink = "https://maps.google.com/?cid=1"
    ws.append(["Barber Max", "facebook.com seulement", "https://www.facebook.com/barbermax", "0490000002", "3 Place Carnot, 84300 Cavaillon, France", "Barbier", 4.2, 18, "", "coiffeur"])
    ws.freeze_panes = "A2"
    wb.save(chemin)


def test_csv_de_l_outil_google_maps(tmp_path):
    chemin = tmp_path / "leads.csv"
    chemin.write_text(CSV_MAPS, encoding="utf-8")
    resultat = lire_prospects(chemin)
    assert resultat.format == FORMAT_MAPS_CSV
    lea, max_, reelle = resultat.prospects
    assert (lea.nom, lea.url, lea.telephone, lea.ville, lea.categorie) == (
        "Salon Léa", "", "04 90 00 00 01", "L'Isle-sur-la-Sorgue", "Coiffeur"
    )
    assert (lea.note_google, lea.nb_avis, lea.lien_maps) == (4.7, 132, "https://maps.google.com/?cid=1")
    assert "sans site" in lea.remarque.lower()
    assert max_.url == "" and "facebook.com/barbermax" in max_.remarque
    assert reelle.url == "https://coiffure-reelle.fr/" and reelle.note_google is None


def test_excel_de_l_outil_google_maps(tmp_path):
    chemin = tmp_path / "leads.xlsx"
    ecrire_xlsx_maps(chemin)
    resultat = lire_prospects(chemin)
    assert resultat.format == FORMAT_MAPS_EXCEL
    lea, max_ = resultat.prospects
    assert (lea.nom, lea.url, lea.ville, lea.note_google, lea.nb_avis) == ("Salon Léa", "", "L'Isle-sur-la-Sorgue", 4.7, 132)
    assert lea.lien_maps == "https://maps.google.com/?cid=1"  # cible de l'hyperlien « Ouvrir »
    assert max_.telephone == "0490000002" and max_.url == "" and "facebook" in max_.remarque


def test_excel_generique(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Nom", "Site", "Ville", "Note", "Avis"])
    wb.active.append(["Garage", "garage.test", "Brest", "4,5", "25"])
    wb.save(tmp_path / "g.xlsx")
    resultat = lire_prospects(tmp_path / "g.xlsx")
    assert resultat.format == FORMAT_GENERIQUE
    (p,) = resultat.prospects
    assert (p.nom, p.url, p.ville, p.note_google, p.nb_avis) == ("Garage", "garage.test", "Brest", 4.5, 25)


def test_excel_illisible(tmp_path):
    from chasseur.importers import ErreurImport

    (tmp_path / "faux.xlsx").write_text("pas un excel")
    with pytest.raises(ErreurImport, match="Excel illisible"):
        lire_prospects(tmp_path / "faux.xlsx")


@pytest.mark.parametrize(
    "entetes, attendu",
    [
        (["name", "status", "website", "phone", "address", "category", "rating", "reviews", "maps_url", "place_id", "query"], FORMAT_MAPS_CSV),
        (["Nom", "Site web", "Page actuelle", "Téléphone", "Adresse", "Catégorie", "Note", "Avis", "Google Maps", "Recherche"], FORMAT_MAPS_EXCEL),
        (["nom", "url"], FORMAT_GENERIQUE),
    ],
)
def test_detection_du_format(entetes, attendu):
    assert detecter_format(entetes) == attendu


@pytest.mark.parametrize(
    "adresse, ville",
    [
        ("12 Rue de la Paix, 75002 Paris, France", "Paris"),
        ("Zone artisanale, 29200 Brest", "Brest"),
        ("Place du Marché, Lourmarin, France", "Lourmarin"),
        ("2 rue du Vide, Lyon", "Lyon"),
        ("Lyon", ""),
        ("", ""),
    ],
)
def test_ville_depuis_adresse(adresse, ville):
    assert ville_depuis_adresse(adresse) == ville
