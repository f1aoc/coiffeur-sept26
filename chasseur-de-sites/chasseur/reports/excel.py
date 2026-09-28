"""Export Excel et CSV de la liste filtrée (§4).

Une ligne par prospect, une colonne par contrôle. Excel : en-têtes figés,
filtres automatiques, couleur par catégorie, onglet « Récapitulatif » (nombre
de prospects par catégorie, par problème et par statut). Les clés API et les
données techniques internes ne sont jamais exportées.
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from datetime import datetime

from sqlmodel import Session

from chasseur.controles import CONTROLES, libelle
from chasseur.db import depot
from chasseur.db.moteur import Stockage
from chasseur.db.tables import STATUTS_COMMERCIAUX, ProspectDB, heure_locale
from chasseur.modeles import Prospect, Resultat
from chasseur.scoring import statut_controles

CONTROLES_EXPORTES = [c for c in CONTROLES if c.id != "url"]
COLONNES_FIXES = [
    "Nom", "Ville", "Catégorie", "Score", "Statut", "Problèmes principaux", "Téléphone", "Site", "Adresse",
    "Activité", "Note Google", "Avis Google", "Analysé le", "Analyse",
]
COLONNES_CAPTURES = ["Capture ordinateur", "Capture téléphone"]
COLONNES = COLONNES_FIXES + [c.libelle for c in CONTROLES_EXPORTES] + COLONNES_CAPTURES

ORDRE_ETATS = ["Cassé", "Obsolète", "Correct", "À revérifier", "Sans site"]
COULEURS = {  # (fond, texte)
    "Cassé": ("FDECEA", "B42318"),
    "Obsolète": ("FFF4E0", "A15C00"),
    "Correct": ("E6F4EA", "1E7A3C"),
}


def _statut_lisible(statut: str) -> str:
    if statut.startswith("KO : "):
        return "Problème : " + ", ".join(libelle(code) for code in statut[5:].split(", "))
    return {"OK": "OK", "non vérifié": "Non vérifié", "n/a": "—"}.get(statut, statut)


def lignes(stockage: Stockage, session: Session, filtres: depot.Filtres) -> list[dict]:
    prospects = depot.prospects_filtres(session, filtres)
    ids = [p.id for p in prospects]
    constats = depot.constats_par_prospect(session, ids)
    noms_scans = {s.id: s.nom for s in depot.scans(session)}
    captures: dict[int, dict] = {i: {} for i in ids}
    for c in depot.captures_par_prospect(session, ids):
        captures[c.prospect_id][c.type] = str(stockage.absolu(c.chemin))
    sortie = []
    for p in prospects:
        r = Resultat(Prospect(), non_verifies=p.non_verifies, sans_objet=p.sans_objet)
        r.constats = constats[p.id]
        statuts = statut_controles(r) if p.analyse else {}
        analyse_le = heure_locale(p.analyse_le)
        ligne = {
            "Nom": p.nom,
            "Ville": p.ville,
            "Catégorie": p.etat or "Non analysé",
            "Score": p.score if p.analyse else None,
            "Statut": p.statut,
            "Problèmes principaux": ", ".join(lib for _, lib, _ in depot.problemes_principaux(constats[p.id])),
            "Téléphone": p.telephone,
            "Site": p.url,
            "Adresse": p.adresse,
            "Activité": p.categorie,
            "Note Google": p.note_google,
            "Avis Google": p.nb_avis,
            "Analysé le": analyse_le.strftime("%d/%m/%Y %H:%M") if analyse_le else "",
            "Analyse": noms_scans.get(p.scan_id, ""),
            "Capture ordinateur": captures[p.id].get("bureau", ""),
            "Capture téléphone": captures[p.id].get("mobile", ""),
        }
        for c in CONTROLES_EXPORTES:
            ligne[c.libelle] = _statut_lisible(statuts[c.id]) if statuts else ""
        sortie.append(ligne)
    return sortie


def exporter_csv(stockage: Stockage, session: Session, filtres: depot.Filtres) -> bytes:
    tampon = io.StringIO(newline="")
    ecrivain = csv.DictWriter(tampon, fieldnames=COLONNES, delimiter=";")
    ecrivain.writeheader()
    for ligne in lignes(stockage, session, filtres):
        ecrivain.writerow({k: ("" if v is None else v) for k, v in ligne.items()})
    return ("﻿" + tampon.getvalue()).encode("utf-8")  # BOM : ouverture directe dans Excel


def exporter_excel(stockage: Stockage, session: Session, filtres: depot.Filtres, description: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    donnees = lignes(stockage, session, filtres)
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Prospects"
    feuille.append(COLONNES)
    entete = PatternFill("solid", fgColor="1F2A3A")
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = entete
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
    feuille.row_dimensions[1].height = 42

    remplissages = {etat: (PatternFill("solid", fgColor=f), Font(bold=True, color=t)) for etat, (f, t) in COULEURS.items()}
    probleme = Font(color="B42318")
    col_categorie = COLONNES.index("Catégorie") + 1
    premiere_col_controle = len(COLONNES_FIXES) + 1
    derniere_col_controle = len(COLONNES_FIXES) + len(CONTROLES_EXPORTES)
    for ligne in donnees:
        feuille.append([ligne[c] for c in COLONNES])
        rang = feuille.max_row
        if ligne["Catégorie"] in remplissages:
            fond, police = remplissages[ligne["Catégorie"]]
            cellule = feuille.cell(rang, col_categorie)
            cellule.fill, cellule.font = fond, police
        for col in range(premiere_col_controle, derniere_col_controle + 1):
            cellule = feuille.cell(rang, col)
            if str(cellule.value or "").startswith("Problème"):
                cellule.font = probleme
    feuille.freeze_panes = "B2"  # en-têtes et colonne « Nom » figés
    feuille.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{max(1, feuille.max_row)}"
    largeurs = {"Nom": 30, "Ville": 16, "Catégorie": 13, "Score": 7, "Statut": 14, "Problèmes principaux": 45,
                "Téléphone": 16, "Site": 32, "Adresse": 38, "Activité": 18, "Note Google": 8, "Avis Google": 8,
                "Analysé le": 16, "Analyse": 18}
    for i, nom in enumerate(COLONNES, start=1):
        feuille.column_dimensions[get_column_letter(i)].width = largeurs.get(nom, 22)

    # Onglet récapitulatif
    recap = classeur.create_sheet("Récapitulatif")
    gras = Font(bold=True)
    recap.append(["Récapitulatif de l'export"])
    recap["A1"].font = Font(bold=True, size=14)
    recap.append([f"Généré le {datetime.now():%d/%m/%Y à %H:%M} · {len(donnees)} prospect(s)" + (f" · {description}" if description else "")])
    recap.append([])

    def tableau(titre: str, compte: list[tuple[str, int]], couleurs: bool = False):
        recap.append([titre, "Prospects"])
        for cellule in recap[recap.max_row]:
            cellule.font = gras
        for nom, n in compte:
            recap.append([nom, n])
            if couleurs and nom in remplissages:
                fond, police = remplissages[nom]
                recap.cell(recap.max_row, 1).fill, recap.cell(recap.max_row, 1).font = fond, police
        recap.append([])

    par_etat = Counter(l["Catégorie"] for l in donnees)
    ordre = ORDRE_ETATS + sorted(set(par_etat) - set(ORDRE_ETATS))
    tableau("Par catégorie", [(e, par_etat[e]) for e in ordre if par_etat[e]], couleurs=True)
    par_probleme = Counter(c.libelle for l in donnees for c in CONTROLES_EXPORTES if l[c.libelle].startswith("Problème"))
    tableau("Par problème", par_probleme.most_common())
    par_statut = Counter(l["Statut"] for l in donnees)
    tableau("Par statut commercial", [(s, par_statut[s]) for s in STATUTS_COMMERCIAUX if par_statut[s]])
    recap.column_dimensions["A"].width = 44
    recap.column_dimensions["B"].width = 12

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


__all__ = ["COLONNES", "exporter_csv", "exporter_excel", "lignes", "ProspectDB"]
