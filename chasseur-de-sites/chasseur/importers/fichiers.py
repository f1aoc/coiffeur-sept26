"""Lecture d'une liste de prospects (CSV ou Excel) et détection de son format.

Formats reconnus automatiquement :
- export CSV de l'outil Google Maps existant (find_no_website.py / Bridge to Leads) :
  name, status, website, phone, address, category, rating, reviews, maps_url, place_id, query ;
- export Excel du même outil : Nom, Site web, Page actuelle, Téléphone, Adresse,
  Catégorie, Note, Avis, Google Maps, Recherche ;
- tout autre CSV/Excel : colonnes nom, url, téléphone, adresse, catégorie, ville,
  note, avis, toutes optionnelles (en-têtes sans casse ni accents, synonymes courants).

L'outil Google Maps ne garde que les entreprises SANS vrai site : `status` vaut
« none » (aucun site) ou « social_or_directory:<hôte> » (page Facebook, Planity…).
Ces lignes sont importées sans URL (état « Sans site »), la page tierce est notée
en remarque. Une ligne « real » (vrai site) serait analysée normalement.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from chasseur.modeles import Prospect

FORMAT_GENERIQUE = "CSV / Excel générique"
FORMAT_MAPS_CSV = "Outil Google Maps (CSV)"
FORMAT_MAPS_EXCEL = "Outil Google Maps (Excel)"
EXTENSIONS_EXCEL = (".xlsx", ".xlsm")


class ErreurImport(ValueError):
    pass


SYNONYMES = {
    "nom": {"nom", "name", "entreprise", "raison sociale", "societe", "commerce"},
    "url": {"url", "site", "site web", "site internet", "website", "web"},
    "telephone": {"telephone", "tel", "phone", "portable", "numero"},
    "adresse": {"adresse", "address", "adresse postale"},
    "categorie": {"categorie", "category", "activite", "secteur", "metier"},
    "ville": {"ville", "city", "commune", "localite"},
    "note_google": {"note", "rating", "note google", "note moyenne"},
    "nb_avis": {"avis", "reviews", "nombre d avis", "nb avis", "nombre avis"},
    "lien_maps": {"maps url", "google maps", "lien maps", "lien google maps"},
}


@dataclass
class Import:
    format: str
    prospects: list[Prospect]
    entetes: list[str] = field(default_factory=list)


def normaliser(texte: str) -> str:
    sans_accents = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return " ".join(sans_accents.lower().replace("_", " ").replace("-", " ").replace("'", " ").split())


def detecter_format(entetes: list[str]) -> str:
    noms = {normaliser(e) for e in entetes if e}
    if {"place id", "maps url", "status", "website"} <= noms:
        return FORMAT_MAPS_CSV
    if {"page actuelle", "google maps", "site web", "recherche"} <= noms:
        return FORMAT_MAPS_EXCEL
    return FORMAT_GENERIQUE


# --- Lecture brute ---------------------------------------------------------------


def _lire_texte(chemin: Path) -> str:
    donnees = chemin.read_bytes()
    try:
        return donnees.decode("utf-8-sig")
    except UnicodeDecodeError:
        return donnees.decode("cp1252")  # CSV exporté par Excel en français


def _en_texte(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def lire_tableau(chemin: Path) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    """En-têtes + lignes (numéro de ligne, {en-tête: valeur}) d'un CSV ou d'un Excel."""
    if chemin.suffix.lower() in EXTENSIONS_EXCEL:
        return _lire_excel(chemin)
    texte = _lire_texte(chemin)
    if not texte.strip():
        return [], []
    try:
        dialecte = csv.Sniffer().sniff(texte.splitlines()[0], delimiters=",;\t")
    except csv.Error:
        dialecte = csv.excel
    lecteur = csv.DictReader(io.StringIO(texte, newline=""), dialect=dialecte)
    entetes = list(lecteur.fieldnames or [])
    lignes = []
    for ligne in lecteur:
        lignes.append((lecteur.line_num, {k: _en_texte(v) for k, v in ligne.items() if k is not None}))
    return entetes, lignes


def _lire_excel(chemin: Path) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    from openpyxl import load_workbook

    try:
        # Pas de mode read_only : il ne donne pas accès aux hyperliens (lien Google Maps sur « Ouvrir »).
        classeur = load_workbook(chemin, data_only=True)
    except Exception as e:
        raise ErreurImport(f"Fichier Excel illisible : {e}") from None
    try:
        feuille = classeur.worksheets[0]
        rangees = feuille.iter_rows()
        entetes = [_en_texte(c.value) for c in next(rangees, ())]
        lignes = []
        for numero, cellules in enumerate(rangees, start=2):
            valeurs = {}
            for entete, cellule in zip(entetes, cellules):
                if entete:
                    lien = cellule.hyperlink.target if cellule.hyperlink is not None else None
                    valeurs[entete] = lien or _en_texte(cellule.value)
            lignes.append((numero, valeurs))
    finally:
        classeur.close()
    return entetes, lignes


# --- Conversion en prospects -----------------------------------------------------

RE_VILLE = re.compile(r"\b\d{5}\s+([^,\d][^,]*)")


def ville_depuis_adresse(adresse: str) -> str:
    """« 12 rue X, 84800 L'Isle-sur-la-Sorgue, France » → « L'Isle-sur-la-Sorgue »."""
    if m := RE_VILLE.search(adresse or ""):
        return m[1].strip()
    parties = [p.strip() for p in (adresse or "").split(",") if p.strip()]
    if parties and parties[-1].lower() in ("france", "belgique", "suisse", "luxembourg"):
        parties = parties[:-1]
    # « 2 rue du Vide, Lyon » : la dernière partie, sans chiffres, est la ville
    if len(parties) >= 2 and not any(ch.isdigit() for ch in parties[-1]):
        return parties[-1]
    return ""


def _nombre(texte: str, type_=float):
    texte = (texte or "").replace(",", ".").strip()
    try:
        return type_(float(texte)) if texte else None
    except ValueError:
        return None


def _correspondance(entetes: list[str]) -> dict[str, str]:
    correspondance: dict[str, str] = {}
    for entete in entetes:
        cle = normaliser(entete or "")
        for champ, noms in SYNONYMES.items():
            if cle in noms and champ not in correspondance:
                correspondance[champ] = entete
    return correspondance


def _prospect(numero: int, valeurs: dict[str, str]) -> Prospect | None:
    if not any(valeurs.values()):
        return None
    p = Prospect(
        ligne=numero,
        nom=valeurs.get("nom", ""),
        url=valeurs.get("url", ""),
        telephone=valeurs.get("telephone", ""),
        adresse=valeurs.get("adresse", ""),
        categorie=valeurs.get("categorie", ""),
        ville=valeurs.get("ville", ""),
        note_google=_nombre(valeurs.get("note_google", "")),
        nb_avis=_nombre(valeurs.get("nb_avis", ""), int),
        lien_maps=valeurs.get("lien_maps", ""),
    )
    p.ville = p.ville or ville_depuis_adresse(p.adresse)
    return p


def _depuis_generique(entetes, lignes, nom_fichier: str) -> list[Prospect]:
    correspondance = _correspondance(entetes)
    if not correspondance:
        raise ErreurImport(
            f"Aucune colonne reconnue dans {nom_fichier} (en-têtes lus : {entetes}). "
            "Colonnes attendues : nom, url, téléphone, adresse, catégorie."
        )
    prospects = []
    for numero, ligne in lignes:
        p = _prospect(numero, {champ: ligne.get(entete, "") for champ, entete in correspondance.items()})
        if p:
            prospects.append(p)
    return prospects


def _depuis_maps(entetes, lignes, excel: bool) -> list[Prospect]:
    par_nom = {normaliser(e): e for e in entetes}

    def col(ligne, *noms):
        for n in noms:
            if n in par_nom:
                return ligne.get(par_nom[n], "")
        return ""

    prospects = []
    for numero, ligne in lignes:
        statut = col(ligne, "status", "site web")
        page = col(ligne, "website", "page actuelle")
        valeurs = {
            "nom": col(ligne, "name", "nom"),
            "telephone": col(ligne, "phone", "telephone"),
            "adresse": col(ligne, "address", "adresse"),
            "categorie": col(ligne, "category", "categorie"),
            "note_google": col(ligne, "rating", "note"),
            "nb_avis": col(ligne, "reviews", "avis"),
            "lien_maps": col(ligne, "maps url", "google maps"),
        }
        p = _prospect(numero, valeurs)
        if not p:
            continue
        s = normaliser(statut)
        if s in ("real", "reel") or (excel and page and s not in ("aucun",) and not s.endswith("seulement")):
            p.url = page
        elif page:
            p.remarque = f"Pas de site propre (page {page} seulement)"
        else:
            p.remarque = "Pas de site (fiche Google Maps sans site web)"
        prospects.append(p)
    return prospects


def lire_prospects(chemin: str | Path) -> Import:
    chemin = Path(chemin)
    if not chemin.is_file():
        raise ErreurImport(f"Fichier introuvable : {chemin}")
    entetes, lignes = lire_tableau(chemin)
    if not entetes:
        return Import(FORMAT_GENERIQUE, [], [])
    format_ = detecter_format(entetes)
    if format_ == FORMAT_GENERIQUE:
        prospects = _depuis_generique(entetes, lignes, chemin.name)
    else:
        prospects = _depuis_maps(entetes, lignes, excel=format_ == FORMAT_MAPS_EXCEL)
    return Import(format_, prospects, entetes)
